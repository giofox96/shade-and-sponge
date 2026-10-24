"""Slide figures for the 10 Nov deck (slides/outline_10nov.md) that are not flowcharts or results: 1600x900 SVG + PNG in
slides/fig_v7/, same style as make_infographics.py. Numbers: notes/topic_decision.md §10, notes/tree_master_plan.md §5b,
databases/barcelona/replacement_rule_check.csv, databases/barcelona/porta_canopy_cover.csv.
Then builds slides/shade_and_sponge_10nov.pptx (one full-slide figure per slide, the spoken script as speaker notes;
needs python-pptx). Run from scripts/: python make_slides_v7.py
"""
import re
import pandas as pd
import make_infographics as mi
from make_infographics import t, para, rect, save, ROOT, INK, MUTED, GRID, HEAT, RAIN, TREE, WARN

mi.OUT = ROOT / "slides/fig_v7"
WHITE = "#FFFFFF"


def card(x, y, w, h, big, small, color):
    s = rect(x, y, w, h, color, rx=12, op=0.08) + f'<rect x="{x}" y="{y}" width="8" height="{h}" fill="{color}"/>'
    s += t(x + 34, y + 78, big, 52, color, "bold")
    return s + para(x + 34, y + 120, small, w - 60, 19)[0]


def fig_context():
    b = card(60, 200, 470, 300, "−56%", "Plane trees in ten years: 43,722 trees, 27.45% of the stock → 12% (press, May 2026)", HEAT)
    b += card(565, 200, 470, 300, "15% · 30%", "Master plan 2017–2037: no species above 15%; tree cover up to 30% of the urban surface", TREE)
    b += card(1070, 200, 470, 300, "1 by 1", "Each tree is replaced when it dies or when its street is rebuilt: thousands of decisions", WARN)
    b += para(60, 600, "This is the context, not the problem: the city has already decided how much of each species. "
              "What follows is what it has not decided.", 1400, 24, fill=INK)[0]
    save("v7_02_context", "The decision is already taken", "Barcelona is replacing its plane trees, one at a time", b,
         "Nació Digital 4 May 2026 and Betevé 7 May 2026 (press); Arbres per viure, Pla director de l'arbrat 2017–2037, targets 1 and 3; "
         "notes/tree_master_plan.md.")


def fig_gate1():
    rc = pd.read_csv(ROOT / "databases/barcelona/replacement_rule_check.csv")
    b = t(60, 190, "ON PAPER · Parcs i Jardins 2004 / 2005", 18, TREE, "bold")
    rows = [("< 3.5 m", "small", "Prunus cerasifera"), ("3.5 – 6 m", "medium", "Melia or Jacaranda"), ("> 6 m", "large", "keep the plane")]
    for k, (w, cls, sp) in enumerate(rows):
        y = 220 + k * 92
        b += rect(60, y, 700, 78, TREE, rx=10, op=0.08) + t(84, y + 47, w, 24, INK, "bold") + t(250, y + 47, sp, 24, TREE, "bold")
        b += t(740, y + 47, "Nou Barris row", 14, MUTED, anchor="end")
    b += para(60, 535, "Criteria: tolerance, street type, tradition, technicians' experience.", 700, 20)[0]
    b += t(60, 600, "No heat. No runoff.", 30, HEAT, "bold")
    b += para(60, 650, "The 2011 guide, still in force, no longer prints the table: only a species catalogue.", 700, 19, fill=MUTED)[0]
    b += t(860, 190, "IN PRACTICE · Porta, trees planted 2007–2026", 18, WARN, "bold")
    for k, cls in enumerate(("small", "medium")):
        g = rc[rc.size_class == cls]
        n, m = g.n.sum(), g[g.in_table].n.sum()
        y = 230 + k * 150
        b += t(860, y + 20, f"{'Narrow' if cls == 'small' else 'Medium'} sidewalks · {n} trees planted", 20, INK, "bold")
        b += rect(860, y + 40, 640, 46, GRID, rx=6) + rect(860, y + 40, 640 * m / n, 46, WARN, rx=6)
        b += t(860 + 640 * m / n + 14, y + 72, f"{100 * m / n:.0f}% match the table", 20, WARN, "bold")
    big = rc[rc.size_class == "large"]
    pl = big[big.genus_sp.str.startswith("Platanus")].n.sum()
    b += t(860, 560, f"Wide sidewalks · {big.n.sum()} trees planted, {pl} of them planes", 20, INK, "bold")
    b += para(860, 600, "Rebuilt streets and exceptions follow other lists and cannot be told apart in the inventory, "
              "so ~90% is an upper bound on non-compliance.", 640, 17, fill=MUTED)[0]
    b += rect(60, 730, 1480, 70, INK, rx=10, op=0.05)
    b += t(84, 774, "A rule exists, it is qualitative, and it is not visible in practice. Current list: asked the city (email sent 24 Oct).", 21, INK, "bold")
    save("v7_03_gate1", "Does the city already decide which tree goes where?", "Yes, on paper. Not with numbers, and not in practice", b,
         "Arbrat viari: selecció d'espècies (Parcs i Jardins 2002, rev. 2004) p. 11; Pla de gestió de l'arbrat viari (2005, rev. 2006) §3.4.3; "
         "Gestió de l'arbrat viari (2011) §11; scripts/replacement_rule_check.py; notes/tree_master_plan.md §5b.")


def fig_trees():
    cc = pd.read_csv(ROOT / "databases/barcelona/porta_canopy_cover.csv").iloc[0]
    cov, plane = 100 * cc.canopy_cover, 5.75                 # plane crowns clipped to Porta (notes/topic_decision.md §10.3)
    x0, y0, W_, H_ = 60, 300, 1400, 90
    s = W_ / 35
    b = rect(x0, y0, W_, H_, GRID, rx=8, op=0.5)
    b += rect(x0, y0, plane * s, H_, HEAT, rx=8) + rect(x0 + plane * s, y0, (cov - plane) * s, H_, TREE, rx=0)
    b += t(x0 + plane * s / 2, y0 + 56, "planes", 18, WHITE, "bold", "middle")
    b += t(x0 + plane * s + (cov - plane) * s / 2, y0 + 56, "all other canopy", 18, WHITE, "bold", "middle")
    for v, lab, c in ((cov, f"Porta today {cov:.1f}%", INK), (30, "City target 30%", TREE)):
        x = x0 + v * s
        b += f'<line x1="{x:.0f}" y1="{y0 - 30}" x2="{x:.0f}" y2="{y0 + H_ + 30}" stroke="{c}" stroke-width="3" stroke-dasharray="{"" if c == INK else "8 5"}"/>'
        b += t(x, y0 - 42, lab, 22, c, "bold", "middle")
    b += t(x0 + (cov + 30) / 2 * s, y0 + 56, "−8.3 ha", 22, TREE, "bold", "middle")
    for k in range(0, 36, 5):
        b += t(x0 + k * s, y0 + H_ + 56, f"{k}%", 14, MUTED, anchor="middle")
    b += card(60, 520, 700, 210, "≈ 29%", "of Porta's canopy is plane trees (5.75% of the area, 832 trees)", HEAT)
    b += card(840, 520, 700, 210, "16% ↔ 20%", "All planes replaced by a Brachychiton-size crown → ~16%; by a Tipuana or Jacaranda-size crown → roughly unchanged", TREE)
    save("v7_04_trees", "Are there enough trees?", "Not in Porta, and the planes are the largest part of what it has", b,
         "ICGC LiDAR 26 Sept 2021, canopy = CHM > 2 m (scripts/porta_canopy_cover.py); plan target 1 (30%); "
         "crown ratios at equal trunk size, notes/topic_decision.md §9.9 and §10.3.")


def fig_problem():
    b = rect(60, 220, 1480, 420, INK, rx=16, op=0.05)
    b += para(110, 320, "When Barcelona replaces a plane tree, the only documented rule for the new species is qualitative "
              "(district, sidewalk width, tradition). No rule quantifies heat or runoff, so thousands of replacements will "
              "change each street's shade and rain interception, and nobody measures by how much.", 1380, 38, lh=1.35, fill=INK, weight="bold")[0]
    for k, (lab, txt, c) in enumerate((("CONTEXT", "plane removal, 15% cap, one tree at a time", MUTED),
                                       ("PROBLEM", "a qualitative rule, not visible in practice", HEAT),
                                       ("GAP", "no quantities for heat or runoff, per position", RAIN))):
        x = 60 + k * 500
        b += t(x, 720, lab, 18, c, "bold") + t(x, 752, txt, 19, INK)
    save("v7_05_problem", "The problem", "Context, problem and gap, kept apart", b, "notes/topic_decision.md §10.1–10.2.")


def fig_gap():
    cols = [("IN PRACTICE", HEAT, ["Rules for which species: catalogue, sidewalk size, 15% cap",
                                   "No number for what a choice does to heat or runoff"]),
            ("IN RESEARCH", RAIN, ["Heat and runoff for fixed scenarios, without species traits (Mannucci et al. 2025; Wu et al. 2024)",
                                   "Optimisers for heat only (Shaamala et al. 2025; Peng et al. 2026)",
                                   "Trait-based recommender with the runoff module switched off (SylvCiT; Nicol et al. 2026)"])]
    b = ""
    for k, (head, c, items) in enumerate(cols):
        x = 60 + k * 760
        b += t(x, 200, head, 20, c, "bold")
        y = 240
        for it in items:
            svg, h = para(x + 30, y + 34, it, 640, 22)
            b += rect(x, y, 720, h + 40, c, rx=10, op=0.07) + svg
            y += h + 60
    b += rect(60, 588, 1480, 112, WARN, rx=10, op=0.08) + t(90, 622, "IN THE DIVERSITY LITERATURE", 18, WARN, "bold")
    b += para(90, 656, "Species caps are rules of thumb (Santamour 1990). Targets far below the current share risk that "
              "“well adapted species may be abandoned for underperforming species” (Kendal et al. 2014, p. 416): "
              "the planes, ~27% vs a 15% cap.", 1420, 20)[0]
    b += rect(60, 718, 1480, 76, TREE, rx=10, op=0.10)
    b += t(90, 765, "This method: trait-based, inside the city's rules, scored on heat and runoff, per position.", 24, TREE, "bold")
    save("v7_07_gap", "The gap", "Rules without numbers, tools without species or runoff, and a cap with no evidence behind it", b,
         "notes/topic_decision.md §7, §9.2, §10.2; Santamour 1990 (METRIA 7) pp. 58-64; Kendal, Dobbs & Lohr 2014 (UFUG 13:411-417); "
         "notes/lit/literature_matrix.csv #139-140.")


def fig_rq():
    b = rect(60, 180, 1480, 190, INK, rx=14, op=0.05)
    b += para(90, 235, "Can a trait-based layout method, inside Barcelona's tree rules, choose replacement species for plane trees "
              "and new planting spots that retain and add more pedestrian cooling and storm-water interception than the city's "
              "own practice?", 1420, 29, lh=1.3, weight="bold")[0]
    sq = [("SQ1 · heat", "Mean UTCI 12–17 h on a hot design day", HEAT),
          ("SQ2 · runoff", "Total and at flood hotspots, 2-yr 60-min storm; larger than the model resolves (Δmin)?", RAIN),
          ("SQ3 · trade-off", "Do the heat-best and runoff-best layouts differ, and in which species?", WARN),
          ("SQ4 · additions", "How much of the gap to 30% can new street trees close, and where?", TREE)]
    for k, (h, txt, c) in enumerate(sq):
        x, y = 60 + (k % 2) * 760, 410 + (k // 2) * 150
        b += rect(x, y, 720, 125, c, rx=10, op=0.08) + t(x + 26, y + 44, h, 22, c, "bold")
        b += para(x + 26, y + 80, txt, 670, 19)[0]
    b += t(60, 760, "Changed after scrutiny: 'reduce' → 'retain' (vs the city's rule, today as reference) · additions co-equal · "
           "'not resolved' is an allowed answer", 18, MUTED)
    save("v7_08_rq", "The research question, scrutinised", "One main question, four sub-questions, each with a metric",
         b, "notes/topic_decision.md §10.5 (scrutiny table) and §10.6.")


def fig_hyp():
    items = [("H1", "Trait-based layouts beat the city's rule on heat and on runoff", "Runoff: a difference below Δmin = 91 m³ is 'not resolved'", TREE),
             ("H2", "The heat-best and runoff-best layouts differ in species", "Tested on the Pareto front (SQ3)", WARN),
             ("Boundary", "At extreme storms, street trees cannot matter", "Was H3; true by construction of the model, so declared, not tested", MUTED)]
    b = ""
    for k, (h, txt, sub, c) in enumerate(items):
        y = 200 + k * 190
        b += rect(60, y, 1480, 160, c, rx=12, op=0.08) + f'<rect x="60" y="{y}" width="8" height="160" fill="{c}"/>'
        b += t(100, y + 70, h, 40, c, "bold") + t(340, y + 66, txt, 28, INK, "bold") + t(340, y + 110, sub, 20, MUTED)
    save("v7_09_hyp", "Hypotheses", "Two to test, one declared limit", b, "notes/topic_decision.md §8, §9.5, §10.5.")


def fig_next():
    steps = [("to 17 Nov", "UTCI with Ladybug on the final layouts · city data check", HEAT),
             ("to 1 Dec", "Replacement-priority score (tree-health proxy) · city screening layer", TREE),
             ("to 8 Dec", "Project tool anywhere in Barcelona · Sant Antoni transfer test", RAIN),
             ("to 18 Dec", "Validation write-up · booklet · final presentation", INK)]
    b = ""
    for k, (when, what, c) in enumerate(steps):
        y = 200 + k * 105
        b += rect(60, y, 1000, 85, c, rx=10, op=0.08) + t(90, y + 52, when, 24, c, "bold") + t(260, y + 52, what, 21, INK)
    b += t(1120, 200, "QUESTIONS FOR YOU", 20, WARN, "bold")
    for k, q in enumerate(("Infrared City access for UTCI?", "Is a within-city transfer (Sant Antoni) enough?",
                           "City-wide screening layer: worth it in the time left?")):
        b += para(1120, 250 + k * 110, q, 420, 22, weight="bold")[0]
    save("v7_15_next", "What comes next", "Porta first, then the whole city", b, "notes/citywide_plan.md §4; notes/methodology.md §10.")


SLIDES = [  # (outline heading number, figure path relative to slides/)
    (1, "fig/01_title.png"), (2, "fig_v7/v7_02_context.png"), (3, "fig_v7/v7_03_gate1.png"), (4, "fig_v7/v7_04_trees.png"),
    (5, "fig_v7/v7_05_problem.png"), (6, "fig_v6/v6_04_crowns.png"), (7, "fig_v7/v7_07_gap.png"), (8, "fig_v7/v7_08_rq.png"),
    (9, "fig_v7/v7_09_hyp.png"), (10, "fig_v7/flow_00_overview.png"), (11, "fig_v7/flow_04_design.png"),
    (12, "fig_v7/res_01_radar.png"), (13, "fig_v7/res_03_scenarios.png"), (14, "fig_v7/flow_05_validation.png"),
    (None, "fig_v7/v7_15_next.png"),
    (None, "fig_v7/flow_02_heat.png"), (None, "fig_v7/flow_03_runoff.png"), (None, "fig_v7/flow_01_screening.png"),
    (None, "fig_v7/res_02_bars.png"), (None, "fig_v6/v6_09_sizing_rule.png")]


def build_pptx():
    from pptx import Presentation
    from pptx.util import Emu
    text = (ROOT / "slides/outline_10nov.md").read_text(encoding="utf-8")
    notes = {int(m.group(1)): m.group(2).strip() for m in
             re.finditer(r"^## (\d+)\..*?\*\*Say \(\d+ words\):\*\*\s*\n\"(.+?)\"\s*$", text, re.S | re.M)}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)       # 16:9
    for n, fig in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(str(ROOT / "slides" / fig), 0, 0, prs.slide_width, prs.slide_height)
        s.notes_slide.notes_text_frame.text = notes.get(n, "Backup." if n is None and "next" not in fig else "")
    out = ROOT / "slides/shade_and_sponge_10nov.pptx"
    prs.save(out)
    assert len(notes) == 14, sorted(notes)
    print("wrote", out.relative_to(ROOT), f"({len(SLIDES)} slides, {len(notes)} with a script)")


if __name__ == "__main__":
    fig_context(); fig_gate1(); fig_trees(); fig_problem(); fig_gap(); fig_rq(); fig_hyp(); fig_next()
    build_pptx()
